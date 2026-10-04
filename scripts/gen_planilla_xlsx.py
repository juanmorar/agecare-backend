#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera PLANILLA_DE_EVALUACION_AVANCE_FASE2.xlsx replicando la planilla
original DUOC (formato .xlsx) a partir del .md.

Hojas:
  - Evaluacion: integrantes + 8 indicadores con nivel de logro (lista
    desplegable), puntaje calculado por formula, total y nota por BUSCARV.
  - Rubrica: descripcion de los 4 niveles por indicador.
  - Escala de notas: tabla puntaje->nota (usada por BUSCARV).
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
OUT = ROOT / 'entrega_fase2' / 'Evidencias Grupales' / 'PLANILLA_DE_EVALUACION_AVANCE_FASE2.xlsx'

# --- Estilos ---
AZUL = PatternFill('solid', fgColor='2E5A88')
GRIS = PatternFill('solid', fgColor='F2F2F2')
VERDE = PatternFill('solid', fgColor='E2EFDA')
AMAR = PatternFill('solid', fgColor='FFF2CC')
BLANCO_BOLD = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(cell):
    cell.fill = AZUL
    cell.font = BLANCO_BOLD
    cell.alignment = CENTER
    cell.border = BORDER


# --- Datos ---
INDICADORES = [
    (1, "Propone ajustes al Proyecto APT considerando dificultades, facilitadores y retroalimentación.", 10),
    (2, "Aplica una metodología que permite el logro de los objetivos propuestos, de acuerdo a los estándares de la disciplina.", 10),
    (3, "Genera evidencias que dan cuenta del avance del Proyecto APT en relación a documentación, programación y almacenamiento de datos, de acuerdo a lo planificado por el equipo y que cumpla con estándares de desarrollo de la industria", 25),
    (4, "Utiliza de manera precisa el lenguaje técnico en los entregables de acuerdo con lo requerido por la disciplina.", 5),
    (5, "Utiliza reglas de redacción, ortografía (literal, puntual, acentual) y las normas para citas y referencias.", 5),
    (6, "Entrega la documentación y evidencias requerida por la asignatura de acuerdo a la estrucutra y nombres solicitados, guardando todas las evidencias de avances en Git", 20),
    (7, "Generan evidencias claras dentro del repositorio del aporte de cada uno de los integrantes del equipo que permitan identificar la equidad en el trabajo y la participación de cada estudiante.", 15),
    (8, "Demuestra un trabajo en equipo en donde todos los miembros del equipo expresan con fluidez el conocimiento del tema expuesto y participan de las actividades planificadas en el proyecto", 10),
]

NIVELES = ["Completamente logrado", "Logrado", "Logro incipiente", "No logrado"]

RUBRICA = {
    1: {
        "Completamente logrado": "Señala los ajustes que realizó o realizará y los justifica considerando las dificultades, facilitadores y retroalimentación del docente.",
        "Logrado": "Señala algunos de los ajustes que realizó o realizará y los justifica considerando las dificultades, facilitadores o retroalimentación del docente.",
        "Logro incipiente": "Señala los ajustes que realizó o realizará, pero no los justifica.",
        "No logrado": "No incluye ajustes ni justifica por qué mantiene su plan inicial.",
    },
    2: {
        "Completamente logrado": "Aplica la metodología definida por el equipo de acuerdo a los estándares de la disciplina, alcanzando los objetivos propuestos para el avance del proyecto.",
        "Logrado": "Aplica la metodología definida de acuerdo a los estándares de la disciplina, pero no se observa el cumplimiento de objetivos propuestos para el avance del proyecto.",
        "Logro incipiente": "Aplica la metodología definida cumpliendo parcialmente con los estándares de la disciplina y con los objetivos propuestos para el avance del proyecto.",
        "No logrado": "Aplica la metodología definida sin cumplir los estándares de la disciplina ni los objetivos propuestos para el avance del proyecto.",
    },
    3: {
        "Completamente logrado": "Presenta evidencias de avance que cumplen con lo planificado por el equipo en relación a documentación, programación y almacenamiento de datos, entregando evidencias que cumplan con estándares de desarrollo que actualmente se encuentran presentes en la industria",
        "Logrado": "Presenta evidencias de avance sin cumplir con lo planificado por el equipo en relación a documentación, programación y almacenamiento de datos, pero si sus entregas evidencian el cumplimiento de estándares de desarrollo que actualmente se encuentran presentes en la industria.",
        "Logro incipiente": "Presenta evidencias de avance que cumplen con lo planificado por el equipo en relación a documentación, programación y almacenamiento de datos, entregando evidencias que no cumplen con estándares de desarrollo que actualmente se encuentran presentes en la industria",
        "No logrado": "Presenta evidencias de avance sin cumplir con lo planificado por el equipo en relación a documentación, programación y almacenamiento de datos y las evidencias no cumplen con estándares de desarrollo que actualmente se encuentran presentes en la industria",
    },
    4: {
        "Completamente logrado": "Utiliza lenguaje técnico de su disciplina en todos los entregables de avance del proyecto.",
        "Logrado": "Utiliza lenguaje técnico de su disciplina en la mayoría de los entregables de avance del proyecto",
        "Logro incipiente": "Utiliza lenguaje técnico de su disciplina en la menos de la mitad de los entregables de avance del proyecto",
        "No logrado": "No utiliza lenguaje técnico de su disciplina en los entregables de avance del proyecto.",
    },
    5: {
        "Completamente logrado": "El texto cumple con las reglas ortografía y de redacción en todos sus apartados y utiliza correctamente todas las normas de citación y referencias.",
        "Logrado": "El texto presenta de 1 a 5 errores de ortografía, redacción o en las citas y referencias del informe.",
        "Logro incipiente": "El texto presenta de 6 a 10 errores de ortografía, redacción o en las citas y referencias del informe.",
        "No logrado": "El texto presenta más de 10 errores de ortografía, redacción o en las citas y referencias del informe.",
    },
    6: {
        "Completamente logrado": "Entrega la documentación y evidencias requeridas por la asignatura de acuerdo a la estrucutra y nombres solicitados, guardando todas las evidencias de avances en Git",
        "Logrado": "Entrega la documentación y evidencias requeridas por la asignatura de acuerdo a la estrucutra y nombres solicitados, guardando algunas de las evidencias de avances en Git",
        "Logro incipiente": "Entrega la documentación y evidencias requeridas por la asignatura sin una la estrucutra y nombres solicitados, guardando algunas de las evidencias de avances en Git",
        "No logrado": "No entrega a través de Git la documentación y evidencias de avance requeridas por la asignatura",
    },
    7: {
        "Completamente logrado": "Generan evidencias dentro del repositorio del proyecto del aporte de cada uno de los integrantes del equipo que permitan identificar la equidad en el trabajo y la participación de cada estudiante",
        "Logrado": "Generan evidencias dentro del repositorio del proyecto del aporte de cada uno de los integrantes del equipo pero la evidencia no demuestra una equidad en la participación de cada estudiante",
        "Logro incipiente": "Generan evidencias dentro del repositorio del proyecto del aporte de cada uno de los integrantes del equipo pero la evidencia no demuestra la participación de algunos(s) participante(s) del equipo",
        "No logrado": "No generan evidencias dentro del repositorio del proyecto del aporte de cada uno de los integrantes del equipo por lo que no se permite identificar la equidad en el trabajo y la participación de cada estudiante",
    },
    8: {
        "Completamente logrado": "Demuestra un trabajo en equipo en donde todos los miembros del equipo expresan con fluidez el conocimiento del tema expuesto y participan de las actividades planificadas en el proyecto",
        "Logrado": "Demuestra un trabajo en equipo de manera parcial, en donde todos los miembros del equipo expresan con fluidez el conocimiento del tema expuesto, pero no todos participan de manera equitativa de las actividades planificadas en el proyecto",
        "Logro incipiente": "Demuestra de manera deficiente el trabajo en equipo en donde los integrantes expresan con poca fluidez y de manera imparcial el conocimiento del tema expuesto y no todos sus miembros participan de manera equitativa de las actividades planificadas en el proyecto",
        "No logrado": "No demuestran un trabajo en equipo, evidenciando en su presentación la nula colaboración de alguno de sus integrantes.",
    },
}

ESCALA = [
    (1.0,0,0.5),(1.1,1,2.5),(1.2,3,4.5),(1.3,5,6.5),(1.4,7,8.5),(1.5,9,10.5),
    (1.6,11,12.5),(1.7,13,14.5),(1.8,15,16.5),(1.9,17,18.5),(2.0,19,20.5),
    (2.1,21,22.5),(2.2,23,24.5),(2.3,25,26.5),(2.4,27,28.5),(2.5,29,30.5),
    (2.6,31,32.5),(2.7,33,34.5),(2.8,35,36.5),(2.9,37,38.5),(3.0,39,40.5),
    (3.1,41,42.5),(3.2,43,44.5),(3.3,45,46.5),(3.4,47,48.5),(3.5,49,50.5),
    (3.6,51,52.5),(3.7,53,54.5),(3.8,55,56.5),(3.9,57,58.5),(4.0,59,60.5),
    (4.1,61,61.5),(4.2,62,63),(4.3,63.5,64.5),(4.4,65,65.5),(4.5,66,67),
    (4.6,67.5,68.5),(4.7,69,69.5),(4.8,70,71),(4.9,71.5,72.5),(5.0,73,73.5),
    (5.1,74,75),(5.2,75.5,76.5),(5.3,77,77.5),(5.4,78,79),(5.5,79.5,80.5),
    (5.6,81,81.5),(5.7,82,83),(5.8,83.5,84.5),(5.9,85,85.5),(6.0,86,87),
    (6.1,87.5,88.5),(6.2,89,89.5),(6.3,90,91),(6.4,91.5,92.5),(6.5,93,93.5),
    (6.6,94,95),(6.7,95.5,96.5),(6.8,97,97.5),(6.9,98,99),(7.0,99.5,100),
]

NIVEL_PCT = {"Completamente logrado":1.0,"Logrado":0.6,"Logro incipiente":0.3,"No logrado":0.0}

wb = Workbook()

# ============ HOJA ESCALA ============
esc = wb.active
esc.title = "Escala de notas"
esc.append(["Nota", "Puntaje desde", "Puntaje hasta"])
for c in esc[1]:
    style_header(c)
for nota, desde, hasta in ESCALA:
    esc.append([nota, desde, hasta])
# Para BUSCARV necesitamos: columna de busqueda (desde) a la izquierda, nota a la derecha.
# Añadimos columnas auxiliares E:F = [desde, nota] ordenado para BUSCARV aproximado.
esc["E1"] = "desde(aux)"; esc["F1"] = "nota(aux)"
style_header(esc["E1"]); style_header(esc["F1"])
for idx, (nota, desde, hasta) in enumerate(ESCALA, start=2):
    esc.cell(row=idx, column=5, value=desde)
    esc.cell(row=idx, column=6, value=nota)
esc.column_dimensions['A'].width = 8
esc.column_dimensions['B'].width = 14
esc.column_dimensions['C'].width = 14
esc.column_dimensions['E'].width = 12
esc.column_dimensions['F'].width = 12
for row in esc.iter_rows(min_row=2, max_row=1+len(ESCALA), max_col=6):
    for c in row:
        c.border = BORDER
        c.alignment = CENTER
n_esc = 1 + len(ESCALA)

# ============ HOJA EVALUACION ============
ev = wb.create_sheet("Evaluacion", 0)
ev["A1"] = "PLANILLA DE EVALUACIÓN — AVANCE FASE 2"
ev["A1"].font = Font(bold=True, size=14, color='2E5A88')
ev.merge_cells("A1:E1")
ev["A2"] = "Proyecto APT · Capstone PTY4614 · DUOC UC — Grupo 1"
ev["A2"].font = Font(italic=True, color='595959')
ev.merge_cells("A2:E2")

# Integrantes
r = 4
ev.cell(row=r, column=1, value="Integrantes (nota grupal)").font = BOLD
r += 1
ev.cell(row=r, column=1, value="N°"); ev.cell(row=r, column=2, value="Integrante"); ev.cell(row=r, column=3, value="Nota")
for c in (ev.cell(row=r, column=1), ev.cell(row=r, column=2), ev.cell(row=r, column=3)):
    style_header(c)
nota_cell = None  # se define abajo
integrantes = ["Javier Cerna Chávez", "Benjamín Camus", "Juan Mora"]
int_start = r + 1
for k, nombre in enumerate(integrantes):
    rr = int_start + k
    ev.cell(row=rr, column=1, value=k+1).alignment = CENTER
    ev.cell(row=rr, column=2, value=nombre).alignment = LEFT
    # la nota se enlaza a la celda de nota total (se setea despues)
    for col in (1,2,3):
        ev.cell(row=rr, column=col).border = BORDER

# Tabla de indicadores
tbl_hdr = int_start + len(integrantes) + 2
ev.cell(row=tbl_hdr-1, column=1, value="Evaluación grupal (8 indicadores)").font = BOLD
headers = ["N°", "Aspecto a evaluar", "Ponderación (%)", "Nivel de logro", "Puntaje"]
for ci, h in enumerate(headers, start=1):
    cell = ev.cell(row=tbl_hdr, column=ci, value=h)
    style_header(cell)

first_data = tbl_hdr + 1
for k, (num, texto, pond) in enumerate(INDICADORES):
    rr = first_data + k
    ev.cell(row=rr, column=1, value=num).alignment = CENTER
    ev.cell(row=rr, column=2, value=texto).alignment = LEFT
    ev.cell(row=rr, column=3, value=pond).alignment = CENTER
    ev.cell(row=rr, column=4, value="Completamente logrado").alignment = CENTER
    # Puntaje = ponderacion * pct(nivel). pct se resuelve con anidado IF.
    nivel_ref = f"D{rr}"
    pond_ref = f"C{rr}"
    formula = (f'=IF({nivel_ref}="Completamente logrado",{pond_ref}*1,'
               f'IF({nivel_ref}="Logrado",{pond_ref}*0.6,'
               f'IF({nivel_ref}="Logro incipiente",{pond_ref}*0.3,0)))')
    ev.cell(row=rr, column=5, value=formula).alignment = CENTER
    for col in range(1,6):
        ev.cell(row=rr, column=col).border = BORDER

last_data = first_data + len(INDICADORES) - 1

# Fila total
total_row = last_data + 1
ev.cell(row=total_row, column=2, value="Total").font = BOLD
ev.cell(row=total_row, column=3, value=f"=SUM(C{first_data}:C{last_data})").font = BOLD
ev.cell(row=total_row, column=3).alignment = CENTER
ev.cell(row=total_row, column=5, value=f"=SUM(E{first_data}:E{last_data})").font = BOLD
ev.cell(row=total_row, column=5).alignment = CENTER
for col in range(1,6):
    ev.cell(row=total_row, column=col).border = BORDER
    ev.cell(row=total_row, column=col).fill = AMAR

# Puntaje total y Nota
pt_row = total_row + 2
ev.cell(row=pt_row, column=2, value="Puntaje total").font = BOLD
ev.cell(row=pt_row, column=3, value=f"=E{total_row}").font = BOLD
ev.cell(row=pt_row, column=3).alignment = CENTER
nota_row = pt_row + 1
ev.cell(row=nota_row, column=2, value="Nota").font = BOLD
# BUSCARV aproximado sobre la escala (col E=desde ordenada asc, col F=nota)
nota_formula = (f"=VLOOKUP(E{total_row},'Escala de notas'!$E$2:$F${n_esc},2,TRUE)")
ev.cell(row=nota_row, column=3, value=nota_formula).font = Font(bold=True, size=13, color='2E5A88')
ev.cell(row=nota_row, column=3).alignment = CENTER
ev.cell(row=pt_row, column=3).fill = VERDE
ev.cell(row=nota_row, column=3).fill = VERDE

# Enlazar nota de integrantes a la nota total
for k in range(len(integrantes)):
    rr = int_start + k
    ev.cell(row=rr, column=3, value=f"=C{nota_row}").alignment = CENTER

# Validacion de datos (lista desplegable de niveles) en D de indicadores
dv = DataValidation(type="list", formula1='"%s"' % ",".join(NIVELES), allow_blank=False)
ev.add_data_validation(dv)
dv.add(f"D{first_data}:D{last_data}")

ev.column_dimensions['A'].width = 5
ev.column_dimensions['B'].width = 70
ev.column_dimensions['C'].width = 15
ev.column_dimensions['D'].width = 24
ev.column_dimensions['E'].width = 10

# ============ HOJA RUBRICA ============
ru = wb.create_sheet("Rubrica")
ru.append(["Indicador", "Ponderación (%)", "Nivel de logro", "Descripción"])
for c in ru[1]:
    style_header(c)
rr = 2
for num, texto, pond in INDICADORES:
    for nivel in NIVELES:
        ru.cell(row=rr, column=1, value=f"Indicador {num}").alignment = CENTER
        ru.cell(row=rr, column=2, value=pond).alignment = CENTER
        ru.cell(row=rr, column=3, value=f"{nivel} ({int(NIVEL_PCT[nivel]*100)}%)").alignment = LEFT
        ru.cell(row=rr, column=4, value=RUBRICA[num][nivel]).alignment = LEFT
        for col in range(1,5):
            ru.cell(row=rr, column=col).border = BORDER
        if nivel == "Completamente logrado":
            for col in range(1,5):
                ru.cell(row=rr, column=col).fill = GRIS
        rr += 1
ru.column_dimensions['A'].width = 12
ru.column_dimensions['B'].width = 15
ru.column_dimensions['C'].width = 26
ru.column_dimensions['D'].width = 90

wb.save(OUT)
print("OK", OUT)
